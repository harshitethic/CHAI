"""Render explicit document specifications in CHAI's isolated project runtime."""
import json,sys,os
from xml.sax.saxutils import escape
spec=json.load(open(sys.argv[1],encoding='utf8'))
kind=spec['format'];path=spec['path'];title=spec.get('title','CHAI document')
paragraphs=spec.get('paragraphs',[]);rows=spec.get('rows',[]);slides=spec.get('slides',[])
os.makedirs(os.path.dirname(path) or '.',exist_ok=True)
if kind=='docx':
 from docx import Document
 from docx.shared import Inches,Pt
 d=Document();d.add_heading(title,0)
 for p in paragraphs:d.add_paragraph(p)
 if rows:
  t=d.add_table(rows=0,cols=max(map(len,rows)));t.style='Light Shading Accent 1'
  for row in rows:
   cells=t.add_row().cells
   for i,value in enumerate(row):cells[i].text=str(value)
 d.save(path)
elif kind=='xlsx':
 from openpyxl import Workbook
 from openpyxl.styles import Font,PatternFill
 from openpyxl.utils import get_column_letter
 w=Workbook();s=w.active;s.title='Data'
 for row in rows:s.append(row)
 s.freeze_panes='A2';s.auto_filter.ref=s.dimensions
 for cell in s[1]:cell.font=Font(bold=True,color='FFFFFF');cell.fill=PatternFill('solid',fgColor='305169')
 for col in s.columns:s.column_dimensions[get_column_letter(col[0].column)].width=min(60,max(12,max(len(str(c.value or '')) for c in col)+3))
 w.save(path)
elif kind=='pptx':
 from pptx import Presentation
 from pptx.util import Inches,Pt
 r=Presentation();r.slide_width=Inches(13.333);r.slide_height=Inches(7.5)
 for item in slides or [{'title':title,'body':paragraphs}]:
  slide=r.slides.add_slide(r.slide_layouts[1]);slide.shapes.title.text=item['title']
  frame=slide.placeholders[1].text_frame;frame.clear()
  for i,text in enumerate(item.get('body',[])):
   p=frame.paragraphs[0] if i==0 else frame.add_paragraph();p.text=text;p.font.size=Pt(24)
 r.save(path)
elif kind=='pdf':
 from reportlab.platypus import SimpleDocTemplate,Paragraph,Spacer,Table,TableStyle
 from reportlab.lib.styles import getSampleStyleSheet
 from reportlab.lib import colors
 styles=getSampleStyleSheet();flow=[Paragraph(escape(title),styles['Title']),Spacer(1,18)]
 for p in paragraphs:flow.extend([Paragraph(escape(p).replace('\n','<br/>'),styles['BodyText']),Spacer(1,10)])
 if rows:
  table=Table([[Paragraph(escape(str(c)),styles['BodyText']) for c in row] for row in rows],repeatRows=1)
  table.setStyle(TableStyle([('BACKGROUND',(0,0),(-1,0),colors.HexColor('#dce6ed')),('GRID',(0,0),(-1,-1),.3,colors.grey),('VALIGN',(0,0),(-1,-1),'TOP'),('BOTTOMPADDING',(0,0),(-1,-1),8)]));flow.append(table)
 SimpleDocTemplate(path).build(flow)
else:raise ValueError('Unsupported format')
print(json.dumps({'path':path,'bytes':os.path.getsize(path),'format':kind}))
